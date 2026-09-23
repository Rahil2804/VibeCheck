import json
from io import BytesIO
from math import cos, radians
from typing import Any

from openpyxl import load_workbook

from backend.models import Coordinates
from backend.snapshot import snapshot_connection

TORONTO_NEIGHBOURHOODS_URL = (
    "https://ckan0.cf.opendata.inter.prod-toronto.ca/dataset/"
    "fc443770-ef0a-4025-9c2c-2cb558bfab00/resource/"
    "0719053b-28b7-48ea-b863-068823a93aaa/download/neighbourhoods-4326.geojson"
)
TORONTO_PROFILE_WORKBOOK_URL = (
    "https://ckan0.cf.opendata.inter.prod-toronto.ca/dataset/"
    "6e19a90f-971c-46b3-852c-0c48c436d1fc/resource/"
    "19d4a806-7385-4889-acf2-256f1e079060/download/"
    "nbhd_2021_census_profile_full_158model.xlsx"
)
PROFILE_FIELDS = {
    "total - age groups of the population - 25% sample data": "population_2021",
    "population density per square kilometre": "population_density",
    "median monthly shelter costs for rented dwellings ($)": "median_renter_shelter_cost",
    "% of tenant households spending 30% or more of its income on shelter costs": (
        "renter_cost_burden_percent"
    ),
}


def normalize_toronto_neighbourhood_profile(
    boundaries: dict[str, Any],
    workbook_bytes: bytes,
    point: Coordinates,
) -> dict[str, Any]:
    neighbourhood = find_neighbourhood(boundaries, point)
    if neighbourhood is None:
        return {}
    profile = extract_profile_values(
        workbook_bytes,
        neighbourhood_id=neighbourhood["id"],
        neighbourhood_name=neighbourhood["name"],
    )
    profile = with_derived_population_density(profile, neighbourhood["geometry"])
    return {
        "neighbourhood_id": neighbourhood["id"],
        "neighbourhood_name": neighbourhood["name"],
        "profile_year": 2021,
        **profile,
    }


def find_neighbourhood(
    boundaries: dict[str, Any],
    point: Coordinates,
) -> dict[str, Any] | None:
    for feature in boundaries.get("features", []):
        geometry = feature.get("geometry") or {}
        if not geometry_contains(geometry, point):
            continue
        properties = feature.get("properties") or {}
        neighbourhood_id = (
            properties.get("AREA_SHORT_CODE")
            or properties.get("AREA_ID")
            or properties.get("_id")
        )
        name = (
            properties.get("AREA_NAME")
            or properties.get("AREA_DESC")
            or properties.get("name")
        )
        if neighbourhood_id is not None and name:
            return {
                "id": str(neighbourhood_id),
                "name": str(name),
                "geometry": geometry,
            }
    return None


def lookup_bundled_toronto_profile(point: Coordinates) -> dict[str, Any]:
    connection = snapshot_connection()
    if connection is None:
        return {}
    try:
        rows = connection.execute(
            """
            SELECT neighbourhood_id, neighbourhood_name, geometry_json,
                   population_density, median_renter_shelter_cost,
                   renter_cost_burden_percent
            FROM toronto_neighbourhood_profiles
            """
        ).fetchall()
    except Exception:
        return {}
    finally:
        connection.close()
    for row in rows:
        geometry = json.loads(row["geometry_json"])
        if not geometry_contains(geometry, point):
            continue
        return {
            "neighbourhood_id": str(row["neighbourhood_id"]),
            "neighbourhood_name": str(row["neighbourhood_name"]),
            "profile_year": 2021,
            "population_density": row["population_density"],
            "median_renter_shelter_cost": row["median_renter_shelter_cost"],
            "renter_cost_burden_percent": row["renter_cost_burden_percent"],
        }
    return {}


def extract_all_profile_values(
    workbook_bytes: bytes,
    neighbourhoods: list[dict[str, str]],
) -> dict[str, dict[str, int | float]]:
    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    output: dict[str, dict[str, int | float]] = {}
    for neighbourhood in neighbourhoods:
        column = _neighbourhood_column(
            rows,
            neighbourhood["id"],
            neighbourhood["name"],
        )
        if column is None:
            output[neighbourhood["id"]] = {}
            continue
        output[neighbourhood["id"]] = _profile_values_from_rows(rows, column)
    return output


def extract_profile_values(
    workbook_bytes: bytes,
    *,
    neighbourhood_id: str,
    neighbourhood_name: str,
) -> dict[str, int | float]:
    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    column = _neighbourhood_column(rows, neighbourhood_id, neighbourhood_name)
    if column is None:
        return {}
    return _profile_values_from_rows(rows, column)


def _profile_values_from_rows(
    rows: list[list[Any]],
    column: int,
) -> dict[str, int | float]:
    output: dict[str, int | float] = {}
    for row in rows:
        for cell in row[: min(column, 12)]:
            label = str(cell or "").strip().lower()
            field = PROFILE_FIELDS.get(label)
            if field is None:
                continue
            value = _number(row[column] if column < len(row) else None)
            if value is not None:
                output[field] = value
    return output


def _neighbourhood_column(
    rows: list[list[Any]],
    neighbourhood_id: str,
    neighbourhood_name: str,
) -> int | None:
    target_name = neighbourhood_name.casefold()
    target_id = neighbourhood_id.casefold()
    for row in rows[:30]:
        for index, value in enumerate(row):
            text = str(value or "").strip().casefold()
            if text == target_name or text == target_id:
                return index
            if target_name and target_name in text:
                return index
            if target_id and f"({target_id})" in text:
                return index
    return None


def geometry_contains(geometry: dict[str, Any], point: Coordinates) -> bool:
    raw = geometry.get("coordinates") or []
    geometry_type = geometry.get("type")
    polygons = raw if geometry_type == "MultiPolygon" else [raw]
    return any(
        polygon
        and _ring_contains(polygon[0], point.lng, point.lat)
        and not any(_ring_contains(hole, point.lng, point.lat) for hole in polygon[1:])
        for polygon in polygons
    )


def with_derived_population_density(
    profile: dict[str, int | float],
    geometry: dict[str, Any],
) -> dict[str, int | float]:
    output = dict(profile)
    population = output.pop("population_2021", None)
    area_km2 = geometry_area_km2(geometry)
    if population is not None and area_km2 > 0:
        output["population_density"] = round(float(population) / area_km2, 1)
    return output


def geometry_area_km2(geometry: dict[str, Any]) -> float:
    raw = geometry.get("coordinates") or []
    polygons = raw if geometry.get("type") == "MultiPolygon" else [raw]
    return sum(_polygon_area_km2(polygon) for polygon in polygons if polygon)


def _polygon_area_km2(polygon: list[list[list[float]]]) -> float:
    if not polygon:
        return 0.0
    shell = abs(_ring_area_km2(polygon[0]))
    holes = sum(abs(_ring_area_km2(ring)) for ring in polygon[1:])
    return max(0.0, shell - holes)


def _ring_area_km2(ring: list[list[float]]) -> float:
    if len(ring) < 3:
        return 0.0
    earth_radius_km = 6371.0088
    mean_latitude = radians(
        sum(float(coordinate[1]) for coordinate in ring) / len(ring)
    )
    projected = [
        (
            earth_radius_km * radians(float(coordinate[0])) * cos(mean_latitude),
            earth_radius_km * radians(float(coordinate[1])),
        )
        for coordinate in ring
    ]
    return (
        sum(
            x1 * y2 - x2 * y1
            for (x1, y1), (x2, y2) in zip(projected, projected[1:] + projected[:1])
        )
        / 2
    )


def _ring_contains(ring: list[list[float]], x: float, y: float) -> bool:
    inside = False
    previous = len(ring) - 1
    for current, coordinate in enumerate(ring):
        x1, y1 = coordinate[:2]
        x2, y2 = ring[previous][:2]
        crosses = (y1 > y) != (y2 > y)
        if crosses and x < (x2 - x1) * (y - y1) / (y2 - y1) + x1:
            inside = not inside
        previous = current
    return inside


def _number(value: Any) -> int | float | None:
    if value in (None, "", "x", "F", ".."):
        return None
    try:
        number = float(str(value).replace(",", "").replace("%", ""))
    except (TypeError, ValueError):
        return None
    return int(number) if number.is_integer() else number
