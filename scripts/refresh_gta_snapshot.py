from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import os
import shutil
import sqlite3
import tempfile
import zipfile
from collections import defaultdict
from datetime import UTC, date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable

import httpx
from openpyxl import load_workbook

from backend.snapshot import SNAPSHOT_SCHEMA_VERSION, sha256_file
from backend.sources.ontario.toronto_profiles import (
    TORONTO_NEIGHBOURHOODS_URL,
    TORONTO_PROFILE_WORKBOOK_URL,
    extract_all_profile_values,
    with_derived_population_density,
)


CMHC_URL = (
    "https://assets.cmhc-schl.gc.ca/sites/cmhc/professional/housing-markets-data-and-"
    "research/housing-data-tables/rental-market/rental-market-report-data-tables/2025/"
    "rmr-toronto-2025-en.xlsx?rev=f107330e-4620-4462-9027-8c6ff1522f1a"
)
CYCLING_URL = (
    "https://ckan0.cf.opendata.inter.prod-toronto.ca/dataset/abbe5ee3-e249-4f86-a219-"
    "f0022eaddcc9/resource/023da9a2-8848-4e10-9cad-e7f9119cd874/download/"
    "cycling-network-4326.geojson"
)
BIKE_SHARE_CATALOGUE_URL = (
    "https://ckan0.cf.opendata.inter.prod-toronto.ca/dataset/2b44db0d-eea9-442d-b038-"
    "79335368ad5a/resource/b69873a1-c180-4ccd-a970-514e434b4971/download/"
    "bike-share-gbfs-general-bikeshare-feed-specification.json"
)
CSD_BOUNDARY_SERVICES = (
    "https://geo.statcan.gc.ca/geo_wa/rest/services/2021/"
    "Digital_boundary_files/MapServer/9/query",
    "https://geo.statcan.gc.ca/geo_wa/rest/services/2021/"
    "Cartographic_boundary_files/MapServer/9/query",
)
CENSUS_PROFILE_DOWNLOAD_URL = (
    "https://www12.statcan.gc.ca/census-recensement/2021/dp-pd/prof/details/"
    "download-telecharger/comp/getFile.cfm?LANG=E&GEONO=005&FILETYPE=CSV"
)
GTFS_FEEDS = {
    "TTC": "https://ckan0.cf.opendata.inter.prod-toronto.ca/dataset/7795b45e-e65a-4465-81fc-c36b9dfff169/resource/cfb6b2b8-6191-41e3-bda1-b175c51148cb/download/opendata_ttc_schedules.zip",
    "GO": "https://assets.metrolinx.com/raw/upload/v1683228856/Documents/Metrolinx/Open%20Data/GO-GTFS.zip",
    "UP": "https://assets.metrolinx.com/raw/upload/v1682367798/Documents/Metrolinx/Open%20Data/UP-GTFS.zip",
    "MiWay": "https://www.miapp.ca/GTFS/google_transit.zip",
    "Brampton Transit": "https://www.arcgis.com/sharing/rest/content/items/a355aabd5a8c490186bdce559c9c75fb/data",
    "YRT": "https://www.yrt.ca/google/google_transit.zip",
    "Durham Region Transit": "https://maps.durham.ca/OpenDataGTFS/GTFS_Durham_TXT.zip",
}
AGENCY_BOUNDS = {
    "TTC": (43.58, 43.86, -79.64, -79.11),
    "GO": (43.20, 44.40, -80.20, -78.50),
    "UP": (43.63, 43.72, -79.64, -79.36),
    "MiWay": (43.48, 43.76, -79.82, -79.54),
    "Brampton Transit": (43.61, 43.86, -80.02, -79.54),
    "YRT": (43.75, 44.25, -79.75, -79.17),
    "Durham Region Transit": (43.72, 44.25, -79.30, -78.50),
}
RENT_UNIT_COLUMNS = {
    "studio": (4, 5),
    "one_bedroom": (8, 9),
    "two_bedroom": (12, 13),
    "three_bedroom_plus": (16, 17),
}
VACANCY_UNIT_COLUMNS = {
    "studio": (4, 5),
    "one_bedroom": (9, 10),
    "two_bedroom": (14, 15),
    "three_bedroom_plus": (19, 20),
}
RENT_GEOGRAPHIES = {
    "3520005": ("Toronto", "City of Toronto"),
    "3521005": ("Mississauga", "Mississauga City"),
    "3521010": ("Brampton", "Brampton City"),
    "3521024": ("Caledon", "Zone 24 - Caledon"),
    "3524002": ("Halton Region (Burlington)", "Halton Region"),
    "3524001": ("Oakville", "Zone 23 - Oakville"),
    "3524009": ("Milton / Halton Hills", "Zone 29 - Milton/Halton Hills"),
    "3524015": ("Milton / Halton Hills", "Zone 29 - Milton/Halton Hills"),
    "3519028": (
        "Richmond Hill / Vaughan / King",
        "Zone 25 - Richmond Hill/Vaughan/King",
    ),
    "3519038": (
        "Richmond Hill / Vaughan / King",
        "Zone 25 - Richmond Hill/Vaughan/King",
    ),
    "3519049": (
        "Richmond Hill / Vaughan / King",
        "Zone 25 - Richmond Hill/Vaughan/King",
    ),
    "3519046": ("Aurora / Newmarket / Whitchurch-Stouffville", "Zone 26 - Aurora"),
    "3519048": ("Aurora / Newmarket / Whitchurch-Stouffville", "Zone 26 - Aurora"),
    "3519044": ("Aurora / Newmarket / Whitchurch-Stouffville", "Zone 26 - Aurora"),
    "3519036": ("Markham", "Zone 27 - Markham"),
    "3518001": ("Pickering / Ajax / Uxbridge", "Zone 28 - Pickering/Ajax/Uxbridge"),
    "3518005": ("Pickering / Ajax / Uxbridge", "Zone 28 - Pickering/Ajax/Uxbridge"),
    "3518029": ("Pickering / Ajax / Uxbridge", "Zone 28 - Pickering/Ajax/Uxbridge"),
    "3518013": ("Durham Region (Oshawa)", "Durham Region"),
    "3518009": ("Durham Region (Whitby)", "Durham Region"),
    "3518017": ("Durham Region (Clarington)", "Durham Region"),
    "3518020": ("Durham Region", "Durham Region"),
}
GTA_CSD_IDS = {
    "3518001",
    "3518005",
    "3518009",
    "3518013",
    "3518017",
    "3518020",
    "3518029",
    "3518039",
    "3519028",
    "3519036",
    "3519038",
    "3519044",
    "3519046",
    "3519048",
    "3519049",
    "3519054",
    "3519070",
    "3520005",
    "3521005",
    "3521010",
    "3521024",
    "3524001",
    "3524002",
    "3524009",
    "3524015",
}
MAX_ARTIFACT_BYTES = 50 * 1024 * 1024


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Build the bundled, versioned GTA data snapshot."
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("backend/snapshots/gta_snapshot.sqlite"),
    )
    parser.add_argument("--timeout", type=float, default=90)
    parser.add_argument(
        "--cmhc-workbook",
        type=Path,
        help="Atomically re-normalize CMHC rows from an already downloaded official workbook.",
    )
    args = parser.parse_args()
    if args.cmhc_workbook:
        replace_cmhc_snapshot(args.output, args.cmhc_workbook)
    else:
        refresh_snapshot(args.output, timeout=args.timeout)
    print(f"Snapshot ready: {args.output}")


def refresh_snapshot(output: Path, *, timeout: float = 90) -> None:
    output = output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    database = _staging_file(output)
    manifest_target = output.with_name("manifest.json")
    manifest_path = _staging_file(manifest_target)
    try:
        resources: list[dict[str, Any]] = []
        with httpx.Client(timeout=timeout, follow_redirects=True) as client:
            downloaded: dict[str, bytes] = {}
            for agency, url in GTFS_FEEDS.items():
                downloaded[agency] = _download(client, url)
                resources.append(
                    _resource_record(f"{agency} GTFS", url, downloaded[agency])
                )
            cmhc = _download(client, CMHC_URL)
            cycling = _download(client, CYCLING_URL)
            bike_catalogue = _download(client, BIKE_SHARE_CATALOGUE_URL)
            toronto_boundaries = _download(client, TORONTO_NEIGHBOURHOODS_URL)
            toronto_profiles = _download(client, TORONTO_PROFILE_WORKBOOK_URL)
            csd_boundaries, csd_boundaries_url = _download_csd_boundaries(client)
            census_profiles = _download(client, CENSUS_PROFILE_DOWNLOAD_URL)
            station_url = _station_information_url(json.loads(bike_catalogue))
            stations = _download(client, station_url)
            resources.extend(
                [
                    _resource_record("CMHC 2025 rental tables", CMHC_URL, cmhc),
                    _resource_record("Toronto cycling network", CYCLING_URL, cycling),
                    _resource_record(
                        "Bike Share GBFS catalogue",
                        BIKE_SHARE_CATALOGUE_URL,
                        bike_catalogue,
                    ),
                    _resource_record(
                        "Bike Share station information", station_url, stations
                    ),
                    _resource_record(
                        "Toronto neighbourhood boundaries",
                        TORONTO_NEIGHBOURHOODS_URL,
                        toronto_boundaries,
                    ),
                    _resource_record(
                        "Toronto 2021 neighbourhood profiles",
                        TORONTO_PROFILE_WORKBOOK_URL,
                        toronto_profiles,
                    ),
                    _resource_record(
                        "Statistics Canada 2021 census subdivision boundaries",
                        csd_boundaries_url,
                        csd_boundaries,
                    ),
                    _resource_record(
                        "Statistics Canada 2021 Census Profile",
                        CENSUS_PROFILE_DOWNLOAD_URL,
                        census_profiles,
                    ),
                ]
            )

        connection = sqlite3.connect(database)
        feed_validity: list[dict[str, str | None]] = []
        try:
            _create_schema(connection)
            for agency, content in downloaded.items():
                normalize_gtfs(connection, agency, content, GTFS_FEEDS[agency])
            normalize_cmhc_workbook(connection, cmhc)
            normalize_cycling_geojson(connection, json.loads(cycling))
            normalize_bike_share(connection, json.loads(stations))
            normalize_toronto_neighbourhood_profiles(
                connection,
                json.loads(toronto_boundaries),
                toronto_profiles,
            )
            normalize_census_subdivisions(
                connection,
                json.loads(csd_boundaries),
                census_profiles,
            )
            connection.execute("ANALYZE")
            connection.commit()
            connection.execute("VACUUM")
            feed_validity = [
                {
                    "agency": row[0],
                    "service_date": row[1],
                    "valid_until": row[2],
                }
                for row in connection.execute(
                    "SELECT agency, service_date, feed_end_date FROM feed_status ORDER BY agency"
                )
            ]
        finally:
            connection.close()

        if database.stat().st_size > MAX_ARTIFACT_BYTES:
            raise ValueError("Normalized artifact exceeds the 50 MB repository budget.")
        snapshot_id = f"gta-{datetime.now(UTC):%Y%m%d}-{sha256_file(database)[:12]}"
        manifest = {
            "schema_version": SNAPSHOT_SCHEMA_VERSION,
            "snapshot_id": snapshot_id,
            "edition": "GTA evidence snapshot 2025/2026",
            "editions": {
                "transit": "Feed-specific service dates in feed_validity",
                "rent": "CMHC 2025 Rental Market Report",
                "cycling": "City of Toronto network current at retrieval",
                "bike_share": "Station information current at retrieval",
                "toronto_profiles": "2021 Census, Toronto 158-neighbourhood model",
                "census_subdivisions": "Statistics Canada 2021 Census Profile",
            },
            "created_at": datetime.now(UTC).isoformat(),
            "artifact": output.name,
            "artifact_bytes": database.stat().st_size,
            "artifact_sha256": sha256_file(database),
            "sources": resources,
            "feed_validity": feed_validity,
            "licences": [
                {
                    "name": "Open Government Licence - Toronto",
                    "url": "https://open.toronto.ca/open-data-license/",
                },
                {
                    "name": "Metrolinx Open Data Licence",
                    "url": "https://www.metrolinx.com/en/about-us/open-data",
                },
                {
                    "name": "Municipal open-data terms",
                    "url": "https://www.ontario.ca/page/open-government-licence-ontario",
                },
                {
                    "name": "CMHC data terms",
                    "url": "https://www.cmhc-schl.gc.ca/about-us/copyright-and-permissions",
                },
                {
                    "name": "Open Government Licence - Canada",
                    "url": "https://open.canada.ca/en/open-government-licence-canada",
                },
            ],
        }
        manifest_path.write_text(
            json.dumps(manifest, indent=2) + "\n", encoding="utf-8"
        )
        _validate_built_snapshot(database)
        os.replace(database, output)
        os.replace(manifest_path, manifest_target)
    finally:
        database.unlink(missing_ok=True)
        manifest_path.unlink(missing_ok=True)


def replace_cmhc_snapshot(output: Path, workbook_path: Path) -> None:
    """Repair the CMHC partition from a previously downloaded official workbook."""
    output = output.resolve()
    workbook_path = workbook_path.resolve()
    manifest_target = output.with_name("manifest.json")
    if not output.exists() or not manifest_target.exists():
        raise FileNotFoundError("An existing snapshot and manifest are required.")
    content = workbook_path.read_bytes()
    database = _staging_file(output)
    manifest_path = _staging_file(manifest_target)
    try:
        shutil.copyfile(output, database)
        manifest = json.loads(manifest_target.read_text(encoding="utf-8"))
        connection = sqlite3.connect(database)
        try:
            connection.execute("DELETE FROM rent_benchmarks")
            normalize_cmhc_workbook(connection, content)
            connection.commit()
            connection.execute("VACUUM")
        finally:
            connection.close()
        resources = [
            resource
            for resource in manifest.get("sources", [])
            if resource.get("name") != "CMHC 2025 rental tables"
        ]
        resources.append(_resource_record("CMHC 2025 rental tables", CMHC_URL, content))
        manifest.update(
            {
                "snapshot_id": (
                    f"gta-{datetime.now(UTC):%Y%m%d}-{sha256_file(database)[:12]}"
                ),
                "created_at": datetime.now(UTC).isoformat(),
                "artifact_bytes": database.stat().st_size,
                "artifact_sha256": sha256_file(database),
                "sources": resources,
            }
        )
        manifest_path.write_text(
            json.dumps(manifest, indent=2) + "\n", encoding="utf-8"
        )
        _validate_built_snapshot(database)
        os.replace(database, output)
        os.replace(manifest_path, manifest_target)
    finally:
        database.unlink(missing_ok=True)
        manifest_path.unlink(missing_ok=True)


def _staging_file(target: Path) -> Path:
    """Create a sibling file so Windows inherits the repository directory ACL."""
    descriptor, name = tempfile.mkstemp(
        prefix=f".{target.name}.",
        suffix=".tmp",
        dir=target.parent,
    )
    os.close(descriptor)
    return Path(name)


def normalize_gtfs(
    connection: sqlite3.Connection,
    agency: str,
    content: bytes,
    source_url: str,
) -> None:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        required = {"stops.txt", "routes.txt", "trips.txt", "stop_times.txt"}
        missing = required - set(archive.namelist())
        if missing:
            raise ValueError(f"{agency} GTFS is missing: {', '.join(sorted(missing))}")
        calendar = _csv_rows(archive, "calendar.txt")
        exceptions = _csv_rows(archive, "calendar_dates.txt")
        service_date = choose_regular_weekday(calendar, exceptions)
        active = active_service_ids(calendar, exceptions, service_date)
        routes = {
            row["route_id"]: {
                "name": row.get("route_short_name")
                or row.get("route_long_name")
                or row["route_id"],
                "type": _int(row.get("route_type"), 3),
            }
            for row in _csv_rows(archive, "routes.txt")
        }
        trips = {
            row["trip_id"]: row
            for row in _csv_rows(archive, "trips.txt")
            if row.get("service_id") in active
        }
        stops = {
            row["stop_id"]: row
            for row in _csv_rows(archive, "stops.txt")
            if _float(row.get("stop_lat")) is not None
            and _float(row.get("stop_lon")) is not None
        }
        counts: dict[tuple[str, str, str], set[str]] = defaultdict(set)
        for row in _csv_rows(archive, "stop_times.txt"):
            trip = trips.get(row.get("trip_id", ""))
            if trip is None or row.get("stop_id") not in stops:
                continue
            seconds = parse_gtfs_time(
                row.get("departure_time") or row.get("arrival_time")
            )
            if seconds is None or not _is_peak(seconds):
                continue
            key = (
                row["stop_id"],
                trip.get("route_id", ""),
                trip.get("direction_id", "0"),
            )
            counts[key].add(row["trip_id"])

        bounds = AGENCY_BOUNDS[agency]
        end_dates = [_parse_gtfs_date(row.get("end_date")) for row in calendar]
        feed_end = max((value for value in end_dates if value), default=service_date)
        connection.execute(
            "INSERT INTO feed_status VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            (
                agency,
                source_url,
                service_date.isoformat(),
                feed_end.isoformat(),
                *bounds,
            ),
        )
        for (stop_id, route_id, direction_id), trip_ids in counts.items():
            stop = stops[stop_id]
            route = routes.get(route_id, {"name": route_id, "type": 3})
            connection.execute(
                """
                INSERT INTO transit_stop_service VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    f"{agency}:{stop_id}",
                    agency,
                    stop_id,
                    stop.get("stop_name") or stop_id,
                    float(stop["stop_lat"]),
                    float(stop["stop_lon"]),
                    route_id,
                    route["name"],
                    direction_id,
                    len(trip_ids) / 4,
                    route["type"],
                    service_date.isoformat(),
                ),
            )


def choose_regular_weekday(
    calendar: list[dict[str, str]],
    exceptions: list[dict[str, str]],
) -> date:
    dates = [
        parsed
        for row in calendar
        for parsed in (
            _parse_gtfs_date(row.get("start_date")),
            _parse_gtfs_date(row.get("end_date")),
        )
        if parsed
    ]
    exception_dates = [
        parsed
        for row in exceptions
        if (parsed := _parse_gtfs_date(row.get("date"))) is not None
    ]
    if dates:
        start, end = min(dates), max(dates)
    elif exception_dates:
        start, end = min(exception_dates), max(exception_dates)
    else:
        raise ValueError("GTFS has no service calendar dates.")
    today = date.today()
    candidates = [
        start + timedelta(days=offset)
        for offset in range((end - start).days + 1)
        if (start + timedelta(days=offset)).weekday() == 1
    ]
    if not candidates:
        raise ValueError("GTFS has no Tuesday in its validity period.")
    ordered = sorted(candidates, key=lambda item: (abs((item - today).days), item))
    for candidate in ordered:
        if active_service_ids(calendar, exceptions, candidate):
            return candidate
    raise ValueError("GTFS has no active regular Tuesday service.")


def active_service_ids(
    calendar: list[dict[str, str]],
    exceptions: list[dict[str, str]],
    service_date: date,
) -> set[str]:
    weekday = service_date.strftime("%A").lower()
    active = {
        row["service_id"]
        for row in calendar
        if row.get(weekday) == "1"
        and (_parse_gtfs_date(row.get("start_date")) or date.min) <= service_date
        and (_parse_gtfs_date(row.get("end_date")) or date.max) >= service_date
    }
    compact = service_date.strftime("%Y%m%d")
    for row in exceptions:
        if row.get("date") != compact:
            continue
        if row.get("exception_type") == "1":
            active.add(row["service_id"])
        elif row.get("exception_type") == "2":
            active.discard(row["service_id"])
    return active


def parse_gtfs_time(value: str | None) -> int | None:
    if not value:
        return None
    try:
        hours, minutes, seconds = (int(part) for part in value.split(":"))
    except (TypeError, ValueError):
        return None
    if hours < 0 or not 0 <= minutes < 60 or not 0 <= seconds < 60:
        return None
    return hours * 3600 + minutes * 60 + seconds


def normalize_cmhc_workbook(connection: sqlite3.Connection, content: bytes) -> None:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        rents = _sheet_rows(workbook["Table 1.1.2"])
        vacancies = _sheet_rows(workbook["Table 1.1.1"])
        for geography_id, (scope, row_hint) in RENT_GEOGRAPHIES.items():
            rent_row = _find_row(rents, row_hint)
            vacancy_row = _find_row(vacancies, row_hint)
            if rent_row is None or vacancy_row is None:
                raise ValueError(f"CMHC workbook lacks the expected {row_hint} row.")
            for unit_size, (value_column, quality_column) in RENT_UNIT_COLUMNS.items():
                vacancy_value_column, vacancy_quality_column = VACANCY_UNIT_COLUMNS[
                    unit_size
                ]
                monthly_rent, suppressed = _cmhc_value(rent_row[value_column - 1])
                vacancy_rate, _vacancy_suppressed = _cmhc_value(
                    vacancy_row[vacancy_value_column - 1]
                )
                connection.execute(
                    """
                    INSERT INTO rent_benchmarks VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        geography_id,
                        scope,
                        unit_size,
                        monthly_rent,
                        vacancy_rate,
                        _text(rent_row[quality_column - 1]),
                        _text(vacancy_row[vacancy_quality_column - 1]),
                        int(suppressed),
                        "CMHC 2025 Rental Market Report",
                        2025,
                        CMHC_URL,
                    ),
                )
    finally:
        workbook.close()


def normalize_cycling_geojson(
    connection: sqlite3.Connection, payload: dict[str, Any]
) -> None:
    features = payload.get("features", [])
    for feature_index, feature in enumerate(features):
        geometry = feature.get("geometry") or {}
        properties = feature.get("properties") or {}
        protected = _is_protected(properties)
        for line_index, line in enumerate(_lines(geometry)):
            for segment_index, (start, end) in enumerate(zip(line, line[1:])):
                if len(start) < 2 or len(end) < 2:
                    continue
                length = _distance_m(start[1], start[0], end[1], end[0])
                if length <= 0:
                    continue
                connection.execute(
                    "INSERT INTO cycling_segments VALUES (?, ?, ?, ?, ?)",
                    (
                        f"{feature_index}:{line_index}:{segment_index}",
                        (start[1] + end[1]) / 2,
                        (start[0] + end[0]) / 2,
                        length,
                        int(protected),
                    ),
                )


def normalize_bike_share(
    connection: sqlite3.Connection, payload: dict[str, Any]
) -> None:
    stations = payload.get("data", {}).get("stations", [])
    for station in stations:
        lat, lng = _float(station.get("lat")), _float(station.get("lon"))
        if lat is None or lng is None:
            continue
        connection.execute(
            "INSERT OR REPLACE INTO bike_share_stations VALUES (?, ?, ?, ?)",
            (str(station.get("station_id")), station.get("name"), lat, lng),
        )


def normalize_toronto_neighbourhood_profiles(
    connection: sqlite3.Connection,
    boundaries: dict[str, Any],
    workbook: bytes,
) -> None:
    neighbourhoods: list[dict[str, str]] = []
    features_by_id: dict[str, dict[str, Any]] = {}
    for feature in boundaries.get("features", []):
        properties = feature.get("properties") or {}
        neighbourhood_id = properties.get("AREA_SHORT_CODE") or properties.get(
            "AREA_ID"
        )
        name = properties.get("AREA_NAME") or properties.get("AREA_DESC")
        geometry = feature.get("geometry") or {}
        if neighbourhood_id is None or not name or not geometry:
            continue
        key = str(neighbourhood_id)
        neighbourhoods.append({"id": key, "name": str(name)})
        features_by_id[key] = geometry
    values_by_id = extract_all_profile_values(workbook, neighbourhoods)
    for neighbourhood in neighbourhoods:
        values = with_derived_population_density(
            values_by_id.get(neighbourhood["id"], {}),
            features_by_id[neighbourhood["id"]],
        )
        connection.execute(
            """
            INSERT INTO toronto_neighbourhood_profiles VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                neighbourhood["id"],
                neighbourhood["name"],
                json.dumps(features_by_id[neighbourhood["id"]], separators=(",", ":")),
                values.get("population_density"),
                values.get("median_renter_shelter_cost"),
                values.get("renter_cost_burden_percent"),
            ),
        )


def normalize_census_subdivisions(
    connection: sqlite3.Connection,
    boundaries: dict[str, Any],
    profile_archive: bytes,
) -> None:
    """Bundle official GTA municipal boundaries and selected neutral Census fields."""
    profiles = _extract_census_profiles(profile_archive)
    inserted: set[str] = set()
    for feature in boundaries.get("features", []):
        properties = feature.get("properties") or {}
        geography_id = str(properties.get("CSDUID") or "")
        geometry = feature.get("geometry") or {}
        if geography_id not in GTA_CSD_IDS or not geometry:
            continue
        values = profiles.get(geography_id, {})
        connection.execute(
            """
            INSERT INTO census_subdivisions VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                geography_id,
                str(properties.get("CSDNAME") or geography_id),
                str(properties.get("CSDTYPE") or ""),
                json.dumps(geometry, separators=(",", ":")),
                values.get("population_density"),
                values.get("median_renter_shelter_cost"),
                values.get("renter_cost_burden_percent"),
            ),
        )
        inserted.add(geography_id)
    missing = sorted(GTA_CSD_IDS - inserted)
    if missing:
        raise ValueError(
            "Statistics Canada boundaries lack GTA CSDs: " + ", ".join(missing)
        )


def _extract_census_profiles(
    content: bytes,
) -> dict[str, dict[str, int | float]]:
    wanted = {
        "population density per square kilometre": "population_density",
        "median monthly shelter costs for rented dwellings ($)": (
            "median_renter_shelter_cost"
        ),
        "% of tenant households spending 30% or more of its income on shelter costs": (
            "renter_cost_burden_percent"
        ),
    }
    output: dict[str, dict[str, int | float]] = {
        geography_id: {} for geography_id in GTA_CSD_IDS
    }
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        candidates = [
            info
            for info in archive.infolist()
            if info.filename.casefold().endswith(".csv")
            and "metadata" not in info.filename.casefold()
        ]
        if not candidates:
            raise ValueError("Statistics Canada Census archive contains no data CSV.")
        member = max(candidates, key=lambda info: info.file_size)
        with archive.open(member) as handle:
            rows = csv.DictReader(
                io.TextIOWrapper(handle, encoding="cp1252", newline="")
            )
            for row in rows:
                dguid = _census_row_value(row, "DGUID")
                geography_id = dguid.removeprefix("2021A0005")
                if geography_id not in GTA_CSD_IDS:
                    continue
                label = " ".join(
                    _census_row_value(
                        row,
                        "CHARACTERISTIC_NAME",
                        "CHARACTERISTIC_NAME_NOM",
                    ).split()
                ).casefold()
                field = wanted.get(label)
                if field is None:
                    continue
                value = _census_number(
                    _census_row_value(row, "C1_COUNT_TOTAL", "CTOTAL", "VALUE")
                )
                if value is not None:
                    output[geography_id][field] = value
    return output


def _census_row_value(row: dict[str, Any], *names: str) -> str:
    normalized = {str(key).strip().casefold(): value for key, value in row.items()}
    for name in names:
        value = normalized.get(name.casefold())
        if value not in (None, ""):
            return str(value).strip()
    return ""


def _census_number(value: str) -> int | float | None:
    if value.casefold() in {"", "x", "f", "..", "n/a"}:
        return None
    try:
        number = float(value.replace(",", "").replace("%", ""))
    except ValueError:
        return None
    return int(number) if number.is_integer() else number


def _create_schema(connection: sqlite3.Connection) -> None:
    connection.executescript(
        """
        CREATE TABLE feed_status (
            agency TEXT PRIMARY KEY, source_url TEXT NOT NULL, service_date TEXT NOT NULL,
            feed_end_date TEXT, min_lat REAL, max_lat REAL, min_lng REAL, max_lng REAL
        );
        CREATE TABLE transit_stop_service (
            stop_key TEXT, agency TEXT, stop_id TEXT, stop_name TEXT, stop_lat REAL,
            stop_lng REAL, route_id TEXT, route_name TEXT, direction_id TEXT,
            departures_per_hour REAL, route_type INTEGER, service_date TEXT
        );
        CREATE INDEX transit_location_idx ON transit_stop_service(stop_lat, stop_lng);
        CREATE TABLE rent_benchmarks (
            geography_id TEXT, geography_name TEXT, unit_size TEXT, monthly_rent INTEGER,
            vacancy_rate REAL, quality_code TEXT, vacancy_quality_code TEXT,
            suppressed INTEGER NOT NULL, edition TEXT, reference_year INTEGER,
            source_url TEXT, PRIMARY KEY (geography_id, unit_size)
        );
        CREATE TABLE cycling_segments (
            id TEXT PRIMARY KEY, center_lat REAL, center_lng REAL, length_m REAL,
            protected INTEGER NOT NULL
        );
        CREATE INDEX cycling_location_idx ON cycling_segments(center_lat, center_lng);
        CREATE TABLE bike_share_stations (
            station_id TEXT PRIMARY KEY, name TEXT, lat REAL, lng REAL
        );
        CREATE INDEX bike_location_idx ON bike_share_stations(lat, lng);
        CREATE TABLE toronto_neighbourhood_profiles (
            neighbourhood_id TEXT PRIMARY KEY, neighbourhood_name TEXT NOT NULL,
            geometry_json TEXT NOT NULL, population_density REAL,
            median_renter_shelter_cost INTEGER, renter_cost_burden_percent REAL
        );
        CREATE TABLE census_subdivisions (
            geography_id TEXT PRIMARY KEY, geography_name TEXT NOT NULL,
            geography_type TEXT, geometry_json TEXT NOT NULL,
            population_density REAL, median_renter_shelter_cost INTEGER,
            renter_cost_burden_percent REAL
        );
        """
    )


def _download_csd_boundaries(client: httpx.Client) -> tuple[bytes, str]:
    where = (
        "CSDUID IN ("
        + ",".join(f"'{geography_id}'" for geography_id in sorted(GTA_CSD_IDS))
        + ")"
    )
    failures: list[str] = []
    for service in CSD_BOUNDARY_SERVICES:
        try:
            response = client.get(
                service,
                params={
                    "where": where,
                    "outFields": "CSDUID,CSDNAME,CSDTYPE,PRUID",
                    "returnGeometry": "true",
                    "outSR": "4326",
                    "f": "geojson",
                },
            )
            response.raise_for_status()
            payload = response.json()
            if not isinstance(payload.get("features"), list) or not payload["features"]:
                raise ValueError("response contained no features")
            return response.content, str(response.request.url)
        except (httpx.HTTPError, json.JSONDecodeError, ValueError) as exc:
            status = getattr(getattr(exc, "response", None), "status_code", None)
            failures.append(
                f"{service}: HTTP {status}"
                if status
                else f"{service}: invalid response"
            )
    raise RuntimeError(
        "Statistics Canada boundary services were unavailable: " + "; ".join(failures)
    )


def _download(client: httpx.Client, url: str) -> bytes:
    response = client.get(url)
    response.raise_for_status()
    if not response.content:
        raise ValueError(f"Downloaded resource is empty: {url}")
    return response.content


def _csv_rows(archive: zipfile.ZipFile, name: str) -> list[dict[str, str]]:
    if name not in archive.namelist():
        return []
    with archive.open(name) as handle:
        wrapper = io.TextIOWrapper(handle, encoding="utf-8-sig", newline="")
        return list(csv.DictReader(wrapper))


def _station_information_url(payload: dict[str, Any]) -> str:
    feeds = payload.get("data", {}).get("en", {}).get("feeds", [])
    for feed in feeds:
        if feed.get("name") == "station_information" and feed.get("url"):
            return str(feed["url"])
    raise ValueError("Bike Share catalogue has no station_information feed.")


def _sheet_rows(sheet: Any) -> list[list[Any]]:
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def _find_row(rows: list[list[Any]], needle: str) -> list[Any] | None:
    normalized = needle.casefold()
    labelled = [
        (str(row[0]).strip().casefold(), row)
        for row in rows
        if row and row[0] is not None
    ]
    for label, row in labelled:
        if label.startswith(normalized):
            return row
    for label, row in labelled:
        if normalized in label:
            return row
    return None


def _cmhc_value(value: Any) -> tuple[int | float | None, bool]:
    if value is None or str(value).strip() in {"", "**", "--", "n/a", "N/A"}:
        return None, True
    if isinstance(value, int | float):
        return value, False
    try:
        number = float(str(value).replace(",", "").replace("%", ""))
    except ValueError:
        return None, True
    return (int(number) if number.is_integer() else number), False


def _lines(geometry: dict[str, Any]) -> Iterable[list[list[float]]]:
    coordinates = geometry.get("coordinates") or []
    if geometry.get("type") == "LineString":
        yield coordinates
    elif geometry.get("type") == "MultiLineString":
        yield from coordinates


def _is_protected(properties: dict[str, Any]) -> bool:
    description = " ".join(str(value).lower() for value in properties.values())
    protected_terms = (
        "cycle track",
        "protected",
        "multi-use trail",
        "multi use trail",
        "off-road",
        "off road",
    )
    return any(term in description for term in protected_terms)


def _is_peak(seconds: int) -> bool:
    return 7 * 3600 <= seconds < 9 * 3600 or 16 * 3600 <= seconds < 18 * 3600


def _parse_gtfs_date(value: str | None) -> date | None:
    if not value:
        return None
    try:
        return datetime.strptime(value, "%Y%m%d").date()
    except ValueError:
        return None


def _distance_m(lat1: float, lng1: float, lat2: float, lng2: float) -> float:
    radius = 6_371_000
    phi1, phi2 = math.radians(lat1), math.radians(lat2)
    delta_phi = math.radians(lat2 - lat1)
    delta_lng = math.radians(lng2 - lng1)
    value = (
        math.sin(delta_phi / 2) ** 2
        + math.cos(phi1) * math.cos(phi2) * math.sin(delta_lng / 2) ** 2
    )
    return radius * 2 * math.atan2(math.sqrt(value), math.sqrt(1 - value))


def _resource_record(name: str, url: str, content: bytes) -> dict[str, Any]:
    return {
        "name": name,
        "url": url,
        "retrieved_at": datetime.now(UTC).isoformat(),
        "sha256": hashlib.sha256(content).hexdigest(),
        "bytes": len(content),
    }


def _validate_built_snapshot(path: Path) -> None:
    connection = sqlite3.connect(path)
    try:
        counts = {
            table: connection.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
            for table in (
                "feed_status",
                "transit_stop_service",
                "rent_benchmarks",
                "cycling_segments",
                "bike_share_stations",
                "toronto_neighbourhood_profiles",
                "census_subdivisions",
            )
        }
    finally:
        connection.close()
    empty = [table for table, count in counts.items() if count == 0]
    if empty:
        raise ValueError("Snapshot validation found empty tables: " + ", ".join(empty))
    if counts["feed_status"] != len(GTFS_FEEDS):
        raise ValueError("Snapshot is missing one or more required transit agencies.")
    if counts["toronto_neighbourhood_profiles"] < 150:
        raise ValueError("Snapshot is missing Toronto neighbourhood profiles.")
    if counts["census_subdivisions"] != len(GTA_CSD_IDS):
        raise ValueError("Snapshot is missing GTA census subdivisions.")


def _float(value: Any) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _int(value: Any, default: int) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def _text(value: Any) -> str | None:
    return str(value).strip() if value not in (None, "") else None


if __name__ == "__main__":
    main()
